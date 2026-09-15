"""
Тесты B6 (WebApp-ручки) — `/api/stock/export`, `/api/debts/export`,
`/api/wh/invoices/export`, `/api/wh/counterparties/export`.

Каждая выгрузка проверяется на маленьком засеянном наборе: непустой валидный
.xlsx с ожидаемыми колонками/строками, роль и (для накладных) диапазон дат
уважаются. FastAPI TestClient, мокаем только границу (verify_init_data,
get_notify_bot) — тот же приём, что в `tests/test_warehouse_api.py`.
"""

from __future__ import annotations

import asyncio
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook


class _FakeBot:
    def __init__(self):
        self.docs = []

    async def send_document(self, chat_id, document, caption=None):
        self.docs.append({"chat_id": chat_id, "caption": caption, "bytes": document.data})


@pytest.fixture
def api(isolated_db, monkeypatch):
    import importlib

    import services.rate_limit as rate_limit
    import services.roles as roles
    import webapp.server as server

    importlib.reload(roles)
    rate_limit.reset()

    db = isolated_db
    ids = {"admin": 1, "boss": 100, "mgr": 200, "mgr2": 201, "guest": 300}
    db.set_role(ids["admin"], "admin_user", "Admin", "admin")
    db.set_role(ids["boss"], "boss_user", "Boss", "boss")
    db.set_role(ids["mgr"], "mgr_user", "Manager", "manager")
    db.set_role(ids["mgr2"], "mgr2_user", "Manager Two", "manager")
    db.set_role(ids["guest"], "guest_user", "Guest", "guest")

    fake_bot = _FakeBot()

    async def _fake_get_bot():
        return fake_bot

    monkeypatch.setattr(server, "get_notify_bot", _fake_get_bot)
    monkeypatch.setattr(
        server,
        "verify_init_data",
        lambda init_data: {"id": int(init_data), "first_name": "U", "username": "u"},
    )
    ids["bot"] = fake_bot
    return TestClient(server.app), db, ids


def _last_xlsx(bot: _FakeBot):
    wb = load_workbook(io.BytesIO(bot.docs[-1]["bytes"]))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    return rows[0], rows[1:]


# ─── /api/stock/export ───────────────────────────────────────────────────────


def _seed_stock(db):
    with db.get_conn() as conn:
        cur = db.get_cursor(conn)
        cur.execute(db.q("INSERT INTO warehouses (name) VALUES (?)"), ("Основной",))
        cur.execute(
            db.q("INSERT INTO products (name, unit, category, created_at) VALUES (?, ?, ?, ?)"),
            ("Болт М8", "шт", "Крепёж", db.now_str()),
        )
        pid = cur.lastrowid
        cur.execute(
            db.q("INSERT INTO stock (product_id, warehouse_id, quantity) VALUES (?, ?, ?)"),
            (pid, 1, 25.0),
        )
        conn.commit()


def test_stock_export_valid_nonempty_xlsx(api):
    client, db, ids = api
    _seed_stock(db)
    r = client.post("/api/stock/export", json={"initData": str(ids["mgr"])})
    assert r.status_code == 200, r.text
    header, data = _last_xlsx(ids["bot"])
    assert header == ("Товар", "Единица", "Категория", "Остаток")
    assert len(data) == 1
    assert data[0][0] == "Болт М8"
    assert data[0][3] == 25.0


def test_stock_export_forbidden_for_guest(api):
    client, _db, ids = api
    r = client.post("/api/stock/export", json={"initData": str(ids["guest"])})
    assert r.status_code == 403


# ─── /api/debts/export ───────────────────────────────────────────────────────


def _seed_debt(db, uid, agent_name="Клиент А"):
    from services import counterparties

    cp = asyncio.run(counterparties.create(agent_name))
    db.set_role(uid, "u", "Manager", "manager")
    oid = db.create_order(uid, "Manager", "")
    db.add_order_item(oid, "Товар", "href", 1, "шт", 1000.0)
    db.update_order_agent(oid, str(cp["counterparty_id"]), agent_name)
    with db.get_conn() as conn:
        cur = db.get_cursor(conn)
        cur.execute(
            db.q("UPDATE orders SET payment_type='credit', currency='USD', due_date='2020-01-01' WHERE id=?"),
            (oid,),
        )
        conn.commit()
    db.update_order_status(oid, "shipped")
    return oid


def test_debts_export_manager_sees_only_own(api):
    client, db, ids = api
    _seed_debt(db, ids["mgr"], "Клиент менеджера")
    _seed_debt(db, ids["mgr2"], "Клиент второго менеджера")

    r = client.post("/api/debts/export", json={"initData": str(ids["mgr"])})
    assert r.status_code == 200, r.text
    _header, data = _last_xlsx(ids["bot"])
    names = [row[2] for row in data]
    assert "Клиент менеджера" in names
    assert "Клиент второго менеджера" not in names


def test_debts_export_boss_sees_all_with_aging_bucket(api):
    client, db, ids = api
    _seed_debt(db, ids["mgr"], "Клиент менеджера")
    r = client.post("/api/debts/export", json={"initData": str(ids["boss"])})
    assert r.status_code == 200, r.text
    header, data = _last_xlsx(ids["bot"])
    assert header[-1] == "Просрочка"
    assert any("Просрочено" in row[-1] for row in data)


# ─── /api/wh/invoices/export ─────────────────────────────────────────────────


def _seed_invoice(client, uid, invoice_date):
    return client.post(
        "/api/wh/invoices/create",
        json={
            "initData": str(uid),
            "type": "incoming",
            "warehouse_id": 1,
            "invoice_date": invoice_date,
            "items": [{"product_id": 1, "quantity": 5, "price_cents": 100}],
        },
    )


def test_invoices_export_respects_date_range(api):
    client, db, ids = api
    _seed_stock(db)
    r1 = _seed_invoice(client, ids["mgr"], "2024-01-10")
    assert r1.status_code == 200, r1.text
    r2 = _seed_invoice(client, ids["mgr"], "2024-06-10")
    assert r2.status_code == 200, r2.text

    r = client.post(
        "/api/wh/invoices/export",
        json={"initData": str(ids["mgr"]), "date_from": "2024-01-01", "date_to": "2024-03-01"},
    )
    assert r.status_code == 200, r.text
    _header, data = _last_xlsx(ids["bot"])
    dates = [row[1] for row in data]
    assert dates == ["2024-01-10"]  # только накладная из диапазона


def test_invoices_export_all_when_no_range(api):
    client, db, ids = api
    _seed_stock(db)
    _seed_invoice(client, ids["mgr"], "2024-01-10")
    _seed_invoice(client, ids["mgr"], "2024-06-10")
    r = client.post("/api/wh/invoices/export", json={"initData": str(ids["mgr"])})
    assert r.status_code == 200, r.text
    _header, data = _last_xlsx(ids["bot"])
    assert len(data) == 2


# ─── /api/wh/counterparties/export ───────────────────────────────────────────


def test_counterparties_export_has_purchases_and_debt(api):
    client, db, ids = api
    _seed_debt(db, ids["mgr"], "Клиент с долгом")
    r = client.post("/api/wh/counterparties/export", json={"initData": str(ids["boss"])})
    assert r.status_code == 200, r.text
    header, data = _last_xlsx(ids["bot"])
    assert header == ("Клиент", "Телефон", "Заказов", "Сумма покупок", "Текущий долг")
    row = next(row for row in data if row[0] == "Клиент с долгом")
    assert row[2] == 1  # один заказ
    assert "USD" in row[3]  # сумма покупок
    assert "USD" in row[4]  # текущий долг
