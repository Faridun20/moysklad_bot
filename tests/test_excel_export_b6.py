"""
Тесты B6 — новые Excel-выгрузки `services/excel_export.py` (Каталог/Долги/
Накладные/Клиенты). Чистые функции dict → bytes; читаем результат обратно
через openpyxl, как и построение (симметрично тестам аналитики, которых для
`build_analytics_xlsx` в проекте не было — см. `tests/test_analytics_parallel.py`
для остального аналитического кода, здесь только сами builder'ы).
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from services import excel_export as xe


def _read(xlsx_bytes: bytes):
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    return rows[0], rows[1:]


def test_build_stock_xlsx_has_expected_columns_and_rows():
    rows = [
        {"name": "Болт М8", "unit": "шт", "category": "Крепёж", "quantity": 12.0},
        {"name": "Кабель ПВ", "unit": "м", "category": None, "quantity": 0.0},
    ]
    header, data = _read(xe.build_stock_xlsx(rows))
    assert header == ("Товар", "Единица", "Категория", "Остаток")
    assert len(data) == 2
    assert data[0][0] == "Болт М8"
    assert data[0][3] == 12.0
    assert data[1][2] == "—"  # категория не задана


def test_build_debts_xlsx_columns_and_source_label():
    rows = [
        {
            "source": "order", "title": "#42", "counterparty": "ООО Ромашка",
            "owner_name": "Иван", "due_date": "2024-01-01", "amount": 100.5,
            "currency": "USD", "bucket_label": "Просрочено >90 дней",
        },
        {
            "source": "machine", "title": "JCB 3CX", "counterparty": "Азиз",
            "owner_name": "Руководство", "due_date": None, "amount": 500,
            "currency": "USD", "bucket_label": "Срок не наступил",
        },
    ]
    header, data = _read(xe.build_debts_xlsx(rows))
    assert header == ("Источник", "№", "Контрагент", "Менеджер", "Срок оплаты", "Долг", "Валюта", "Просрочка")
    assert data[0][0] == "Заказ"
    assert data[1][0] == "Техника"
    assert data[1][4] == "—"  # без срока


def test_build_invoices_xlsx_columns():
    rows = [
        {
            "number": "IN-1", "date": "2024-05-01", "type_label": "Приход",
            "counterparty": "Поставщик", "amount": 1000, "currency": "USD",
            "status_label": "Проведена",
        },
    ]
    header, data = _read(xe.build_invoices_xlsx(rows))
    assert header == ("Номер", "Дата", "Тип", "Контрагент", "Сумма", "Валюта", "Статус")
    assert data[0][0] == "IN-1"
    assert data[0][2] == "Приход"


def test_build_counterparties_xlsx_multi_currency_join():
    rows = [
        {
            "name": "ООО Ромашка", "phone": "+998901234567", "orders_count": 3,
            "purchases": [{"currency": "USD", "total": 1500.0}, {"currency": "UZS", "total": 200000.0}],
            "debts": [{"currency": "USD", "total": 300.0}],
        },
        {"name": "Без долгов", "phone": None, "orders_count": 1, "purchases": [], "debts": []},
    ]
    header, data = _read(xe.build_counterparties_xlsx(rows))
    assert header == ("Клиент", "Телефон", "Заказов", "Сумма покупок", "Текущий долг")
    assert "USD" in data[0][3] and "UZS" in data[0][3]
    assert data[1][4] == "—"


def test_formula_injection_is_escaped_in_all_builders():
    """Контрагент/товар вписан менеджером — «=HYPERLINK(...)» не должен стать
    формулой (см. `_text`/`_FORMULA_LEADERS`, тот же риск, что у аналитики)."""
    evil = "=HYPERLINK(\"http://evil\",\"click\")"
    stock_rows = [{"name": evil, "unit": "шт", "category": "", "quantity": 1}]
    _, data = _read(xe.build_stock_xlsx(stock_rows))
    assert data[0][0].startswith("'=")

    cp_rows = [{"name": evil, "phone": None, "orders_count": 0, "purchases": [], "debts": []}]
    _, data = _read(xe.build_counterparties_xlsx(cp_rows))
    assert data[0][0].startswith("'=")
