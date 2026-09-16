"""Юридические документы: пропись, график платежей, сборка контекста, PDF.

Пропись и график проверяются подробно: это денежные поля подписываемого
документа. Сам рендер — по одному на каждый из трёх видов, он пропускается
там, где нет LibreOffice (в CI его нет, в образе есть).
"""

import os
import shutil
from datetime import date
from decimal import Decimal

import pytest

from services import legal_docs as ld
from services.numerals import amount_in_words, ru_number, uz_number

HAS_SOFFICE = shutil.which("soffice") is not None


# ─── Пропись ──────────────────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "n,expected",
    [
        (0, "ноль"), (1, "один"), (2, "два"), (11, "одиннадцать"),
        (14, "четырнадцать"), (21, "двадцать один"), (100, "сто"),
        (101, "сто один"), (200, "двести"), (999, "девятьсот девяносто девять"),
        # Тысяча женского рода, миллион мужского — классическая ошибка прописи.
        (1000, "одна тысяча"), (2000, "две тысячи"), (5000, "пять тысяч"),
        (1_000_000, "один миллион"), (2_000_000, "два миллиона"),
        (5_000_000, "пять миллионов"),
        # 11-14 берут форму «тысяч», а не «тысячи» — исключение из правила.
        (11_000, "одиннадцать тысяч"), (12_000, "двенадцать тысяч"),
        (21_000, "двадцать одна тысяча"), (22_000, "двадцать две тысячи"),
        (25_000, "двадцать пять тысяч"),
        (1234, "одна тысяча двести тридцать четыре"),
    ],
)
def test_ru_number(n, expected):
    assert ru_number(n) == expected


@pytest.mark.parametrize(
    "n,expected",
    [
        (0, "нол"), (1, "бир"), (5, "беш"), (10, "ўн"), (11, "ўн бир"),
        (21, "йигирма бир"), (90, "тўқсон"),
        # «юз» и «минг» без «бир» — так говорят.
        (100, "юз"), (101, "юз бир"), (200, "икки юз"),
        (1000, "минг"), (1001, "минг бир"), (2000, "икки минг"),
        (25_000, "йигирма беш минг"),
        (1234, "минг икки юз ўттиз тўрт"),
        (1_000_000, "бир миллион"),
    ],
)
def test_uz_number(n, expected):
    assert uz_number(n) == expected


def test_amount_in_words_appends_cents_as_fraction():
    """Копейки — дробью «50/100»: принятая в договорах форма, не требующая
    согласования с валютой. У целой суммы хвоста нет."""
    assert amount_in_words(Decimal("25000"), "ru") == "двадцать пять тысяч"
    assert amount_in_words(Decimal("25000.50"), "ru") == "двадцать пять тысяч 50/100"
    assert amount_in_words(Decimal("25000.05"), "uz") == "йигирма беш минг 05/100"


def test_amount_in_words_rejects_unknown_language():
    with pytest.raises(ValueError):
        amount_in_words(Decimal("1"), "en")


# ─── График платежей ──────────────────────────────────────────────────────────


def test_schedule_sums_exactly_to_total():
    """Сумма графика обязана сойтись с суммой договора до сума.

    25 000 000 / 6 не делится нацело: остаток от деления уходит в ПОСЛЕДНИЙ
    платёж. Клиент сложит колонку — она должна дать ровно сумму договора.
    """
    rows = ld.build_schedule(2_500_000_000, 6, date(2026, 9, 12))
    assert len(rows) == 6
    total = sum(int(r["amount"].replace(" ", "")) for r in rows)
    assert total == 25_000_000
    assert rows[-1]["balance"] == "0"
    # Все платежи равны, кроме последнего — он добирает остаток. Тийинов нет.
    assert rows[0]["amount"] == "4 166 666"
    assert rows[-1]["amount"] == "4 166 670"
    assert rows[0]["balance"] == "20 833 334"


def test_schedule_dates_use_month_arithmetic():
    """31 января + 1 месяц = 28 февраля, а не 3 марта: relativedelta, не timedelta."""
    rows = ld.build_schedule(300_000, 3, date(2026, 1, 31))
    assert [r["date"] for r in rows] == ["28.02.2026", "31.03.2026", "30.04.2026"]


def test_schedule_single_payment_is_whole_sum():
    rows = ld.build_schedule(2_500_000_000, 1, date(2026, 9, 12))
    assert len(rows) == 1
    assert rows[0]["amount"] == "25 000 000"
    assert rows[0]["balance"] == "0"


@pytest.mark.parametrize("count,total", [(0, 1000), (-1, 1000), (2, 0), (2, -5), (2, 150)])
def test_schedule_rejects_nonsense(count, total):
    """В том числе тийины (150 копеек = 1,50 сума): сумы в расписке целые."""
    with pytest.raises(ld.DocumentError):
        ld.build_schedule(total, count, date(2026, 9, 12))


# ─── Контекст ─────────────────────────────────────────────────────────────────


# Реквизиты кредитора, которые печатает расписка: название, ИНН, адрес.
# Подписанта кредитора (должность, «в лице …») в бланке больше нет.
FULL_CREDITOR = {
    "name": "FARID IMPEKS LLC", "tin": "309876543", "address": "г. Ташкент, ул. Амира Темура, 107Б",
}


def _ctx(**over):
    base = dict(
        doc_type="raspiska_ru_uz", city="Ташкент",
        creditor=dict(FULL_CREDITOR),
        product_name="Экскаватор JCB 3CX", total_cents=2_500_000_000,
        start_date=date(2026, 9, 12), term_months=6, installments_count=6,
        debtor_full_name="Иванов Иван Иванович",
    )
    base.update(over)
    return ld.build_context(**base)


def test_all_document_types_are_the_handwritten_lawyer_form():
    """Три вида одного бланка: RU+UZ, рус., ўзб. Старые ключи сохранены — на
    них ссылаются записи generated_documents в проде."""
    assert list(ld.TEMPLATES) == ["raspiska_ru_uz", "raspiska_ru", "tilxat_uz"]
    assert set(ld.TEMPLATES) == ld.HANDWRITTEN_TYPES
    assert ld.DOCUMENT_CURRENCY == "UZS"


def _template_text(doc_type) -> str:
    import re
    import zipfile

    xml = zipfile.ZipFile(ld.template_path(doc_type)).read("word/document.xml").decode()
    return re.sub(r"<[^>]+>", "", xml)


@pytest.mark.parametrize("doc_type", list(ld.TEMPLATES))
def test_context_covers_every_template_placeholder(doc_type):
    """Ключи контекста обязаны покрывать плейсхолдеры шаблона.

    Недостающий ключ docxtpl отрисует ПУСТОТОЙ — документ уйдёт клиенту с
    дырой на месте суммы или реквизитов, и ни одна проверка этого не заметит.
    """
    import re

    placeholders = {
        m for m in re.findall(r"\{\{\s*([\w.]+)\s*\}\}", _template_text(doc_type))
        if not m.startswith("item.")
    }
    assert placeholders, "в шаблоне нет ни одного плейсхолдера — собран не тот файл"
    missing = placeholders - set(_ctx(doc_type=doc_type))
    assert not missing, f"{doc_type}: контекст не отдаёт {sorted(missing)}"


def test_amount_words_in_both_languages():
    """Сумма печатается цифрами и прописью, пропись — на языке своей части."""
    ctx = _ctx()
    assert ctx["total_amount"] == "25 000 000"
    assert ctx["total_amount_words_ru"] == "двадцать пять миллионов"
    assert ctx["total_amount_words_uz"] == "йигирма беш миллион"


def test_single_payment_gives_one_row():
    ctx = _ctx(installments_count=1)
    assert len(ctx["schedule"]) == 1
    assert ctx["schedule"][0]["amount"] == "25 000 000"


def test_installments_longer_than_term_rejected():
    """График, который кончается позже срока расписки, противоречит сам себе."""
    with pytest.raises(ld.DocumentError, match="позже срока"):
        _ctx(term_months=3, installments_count=6)


@pytest.mark.parametrize(
    "over,msg",
    [
        ({"doc_type": "нет-такого"}, "тип документа"),
        ({"installments_count": 0}, "больше нуля"),
        ({"term_months": 0}, "не меньше месяца"),
        ({"total_cents": 150}, "тийинов"),
    ],
)
def test_context_validation(over, msg):
    with pytest.raises(ld.DocumentError, match=msg):
        _ctx(**over)


def test_end_date_is_start_plus_term():
    assert _ctx(term_months=6)["end_date"] == "12.03.2027"


def test_requires_only_tin_and_address_and_says_where_to_fill_them():
    """Без ИНН и адреса кредитора первая фраза расписки ушла бы с дырой — отказ
    называет поля поимённо и место, где их заполнить."""
    with pytest.raises(ld.DocumentError) as e:
        _ctx(creditor={"name": "X"})
    assert str(e.value) == (
        "Заполните ИНН и юридический адрес в Настройки → Реквизиты компании — "
        "без этого не выписать расписку"
    )
    with pytest.raises(ld.DocumentError, match="^Заполните ИНН в Настройки → Реквизиты компании"):
        _ctx(creditor={"name": "X", "address": "Ташкент"})


@pytest.mark.parametrize("doc_type", list(ld.TEMPLATES))
def test_creditor_signatory_is_never_asked(doc_type):
    """Жалоба владельца: расписку пишет физлицо, а форма требовала «должность
    подписанта». Ни один вид не спрашивает должность, представителя, основание
    или доверенность — и не кладёт их в контекст."""
    ctx = _ctx(doc_type=doc_type, creditor=dict(FULL_CREDITOR))
    assert not [k for k in ctx if "position" in k or "representative" in k or "basis" in k or "poa" in k]
    assert {key for key, _setting in ld._REQUIRED} == {"tin", "address"}


def test_uzbek_city_falls_back_to_city():
    assert _ctx()["city_uz"] == "Ташкент"
    assert _ctx(creditor={**FULL_CREDITOR, "city_uz": "Тошкент"})["city_uz"] == "Тошкент"


# ─── Шаблоны: три вида одного бланка юриста ───────────────────────────────────


def _docx_paragraphs(path) -> list[str]:
    from docx import Document

    doc = Document(str(path))
    rows = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            rows.extend(cell.text for cell in row.cells)
    return rows


@pytest.mark.parametrize("doc_type", list(ld.TEMPLATES))
def test_committed_templates_match_the_lawyer_source(doc_type, tmp_path):
    """Шаблоны — результат `scripts/build_raspiska_ru_uz` над бланком юриста.

    .docx двоичный: правку в нём не видно в диффе и не отревьюить, а правка
    руками разошлась бы со скриптом и молча пропала при следующей сборке
    (сравниваются тексты, не байты: zip несёт таймстемпы). Заодно: ни одной
    [скобки] бланка не осталось.
    """
    from scripts import build_raspiska_ru_uz as gen

    assert set(gen.VARIANTS) == set(ld.TEMPLATES)
    assert f"{doc_type}.docx" == ld.TEMPLATES[doc_type][0]
    fresh = tmp_path / f"{doc_type}.docx"
    gen.build(doc_type).save(str(fresh))
    committed = _docx_paragraphs(ld.template_path(doc_type))
    assert _docx_paragraphs(fresh) == committed, (
        f"{doc_type}: шаблон разошёлся со скриптом — пересоберите `python -m scripts.build_raspiska_ru_uz`"
    )
    assert not [t for t in committed if "[" in t]


# Что убрано из бланка юриста (решение владельца 2026-09) — ровно это и ничего
# больше. Фрагменты — в виде старых шаблонов (с переменными docxtpl).
_REMOVED_FRAGMENTS = (
    ", в лице {{ creditor_representative_gen }}, действующего на основании {{ creditor_basis_ru }},",
    " номидан {{ creditor_position_uz }} {{ creditor_representative }}, {{ creditor_basis_uz }} асосида иш юритувчи",
)
_REMOVED_LINES = (
    "В лице: ______________________ / {{ creditor_representative }} /",
    "(подпись)",
    "Должность: {{ creditor_position }}, действует на основании {{ creditor_basis_ru }}",
    "М.П.",
    "Номидан: ______________________ / {{ creditor_representative }} /",
    "(имзо)",
    "Лавозими: {{ creditor_position_uz }}, {{ creditor_basis_uz }} асосида иш юритади",
    "М.Ў. (муҳр ўрни)",
)


@pytest.mark.parametrize("doc_type", list(ld.TEMPLATES))
def test_lawyer_text_is_unchanged_except_the_creditor_signatory(doc_type):
    """Весь остальной текст юриста — байт в байт прежний.

    Эталон — абзацы шаблонов до правки (`tests/data/raspiska_text_before_2026_09.json`).
    Из него вычитаются ТОЛЬКО фраза «в лице …, действующего на основании …» (узб.
    «номидан … асосида иш юритувчи») и строки подписанта кредитора; всё прочее —
    пени, статьи ГК, график, подписи должника и свидетеля — обязано совпасть.
    """
    import json
    from pathlib import Path

    before = json.loads(
        (Path(__file__).parent / "data" / "raspiska_text_before_2026_09.json").read_text("utf-8")
    )[doc_type]
    expected = []
    for text in before:
        if text in _REMOVED_LINES:
            continue
        for fragment in _REMOVED_FRAGMENTS:
            text = text.replace(fragment, "")
        expected.append(text)
    assert _docx_paragraphs(ld.template_path(doc_type)) == expected
    removed = [t for t in before if t in _REMOVED_LINES or any(f in t for f in _REMOVED_FRAGMENTS)]
    parts = ld.TEMPLATES[doc_type][1]
    assert len(removed) == 5 * len(parts), "убрано больше или меньше, чем решено"


def test_single_language_templates_hold_only_their_part():
    """Одноязычный вид — ровно своя часть бланка: без текста другого языка и
    без разрыва страницы (он дал бы пустую страницу в PDF)."""
    from docx import Document
    from docx.oxml.ns import qn

    both = _template_text("raspiska_ru_uz")
    ru = _template_text("raspiska_ru").strip()
    uz = _template_text("tilxat_uz").strip()
    assert "РАСПИСКА" in both and "ТИЛХАТ" in both
    assert ru.startswith("РАСПИСКА") and "ТИЛХАТ" not in ru and "Қарздор" not in ru
    assert uz.startswith("ТИЛХАТ") and "РАСПИСКА" not in uz and "Должник" not in uz
    for doc_type in ("raspiska_ru", "tilxat_uz"):
        body = Document(str(ld.template_path(doc_type))).element.body
        assert not [br for br in body.iter(qn("w:br")) if br.get(qn("w:type")) == "page"], doc_type
    assert len(Document(str(ld.template_path("raspiska_ru_uz"))).tables) == 2
    assert len(Document(str(ld.template_path("raspiska_ru"))).tables) == 1
    assert len(Document(str(ld.template_path("tilxat_uz"))).tables) == 1


def test_language_clause_is_edited_only_in_single_language_templates():
    """Единственное отступление от текста юриста: в одноязычном документе нет
    второго языка, на который ссылается пункт 8. В двуязычном — дословно."""
    from scripts import build_raspiska_ru_uz as gen

    both = " ".join(_docx_paragraphs(ld.template_path("raspiska_ru_uz")))
    ru = _docx_paragraphs(ld.template_path("raspiska_ru"))
    uz = _docx_paragraphs(ld.template_path("tilxat_uz"))
    for lang, (old, new) in gen.LANGUAGE_CLAUSE.items():
        assert old in both and new not in both, lang
    assert "8. Настоящая расписка составлена в двух экземплярах, имеющих одинаковую юридическую силу." in ru
    assert "8. Ушбу тилхат бир хил юридик кучга эга бўлган икки нусхада тузилди." in uz
    assert not [t for t in ru + uz if "узбекском" in t or "ўзбек тил" in t]


def test_template_keeps_handwritten_lines_and_layout():
    """Личные данные Должника остаются чертой «от руки»; сумма печатается;
    город и дата — одной строкой через табуляцию (в бланке пробелы, и дата
    уезжала на вторую строку); заголовок графика не отрывается от таблицы;
    строки графика не рвутся, cantSplit — первым в w:trPr."""
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(str(ld.template_path("raspiska_ru_uz")))
    texts = [p.text for p in doc.paragraphs]
    assert sum(t.startswith("Я, ____") for t in texts) == 1
    assert sum(t.startswith("Мен, ____") for t in texts) == 1
    assert sum("Паспорт (серия, номер): ____" in t for t in texts) == 1
    assert "г. {{ city }}\t«___» _______________ 20__ г." in texts
    assert "{{ city_uz }} ш.\t«___» _______________ 20__ й." in texts
    assert sum("общей стоимостью {{ total_amount }} ({{ total_amount_words_ru }}) сум" in t for t in texts) == 1
    assert sum("умумий қиймати {{ total_amount }} ({{ total_amount_words_uz }}) сўм" in t for t in texts) == 1
    headings = [p for p in doc.paragraphs if p.text in ("2. График платежей:", "2. Тўлов жадвали:")]
    assert len(headings) == 2 and all(p.paragraph_format.keep_with_next for p in headings)
    for table in doc.tables:
        cells = [[c.text for c in row.cells] for row in table.rows]
        assert cells[1][0] == "{%tr for item in schedule %}" and cells[3][0] == "{%tr endfor %}"
        assert cells[2] == ["{{ item.number }}", "{{ item.date }}", "{{ item.amount }}", "{{ item.balance }}"]
        for row in table.rows:
            assert row._tr.trPr[0].tag == qn("w:cantSplit")
        assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None


# ─── Рендер ───────────────────────────────────────────────────────────────────


# pypdf выносит «қ», «ҳ», «ғ» из шрифта LibreOffice в конец строки («умумий
# иймати … қ»): глифы без ToUnicode. Сравниваем узбекский текст без этих букв —
# и в PDF, и в ожидаемой фразе, — остальные буквы стоят на местах.
_UZ_DETACHED = str.maketrans("", "", "қҳғҚҲҒ")


def _pdf_text(pdf) -> tuple[str, list[str]]:
    from pypdf import PdfReader

    pages = [p.extract_text() for p in PdfReader(str(pdf)).pages]
    # Пробелы схлопываем: LibreOffice переносит строки по ширине страницы, и
    # «Экскаватор JCB 3CX» может приехать разорванным. Проверяем содержание,
    # а не вёрстку — иначе тест ломается от правки любого поля выше по тексту.
    return " ".join(" ".join(pages).translate(_UZ_DETACHED).split()), pages


@pytest.mark.skipif(not HAS_SOFFICE, reason="нет LibreOffice (в образе он есть)")
@pytest.mark.parametrize("doc_type", list(ld.TEMPLATES))
def test_render_pdf(doc_type, tmp_path):
    import asyncio

    ctx = _ctx(doc_type=doc_type, creditor={**FULL_CREDITOR, "city_uz": "Тошкент"},
               term_months=12, installments_count=12)
    pdf = asyncio.run(ld.render_pdf(doc_type, ctx, tmp_path))
    assert pdf.is_file() and pdf.name.startswith(doc_type)
    text, pages = _pdf_text(pdf)
    parts = ld.TEMPLATES[doc_type][1]

    # Ни один плейсхолдер и ни одна скобка бланка не доехали до подписи.
    assert "{{" not in text and "{%" not in text and "[" not in text
    # ФИО должника в документ не печатается: его пишут от руки.
    assert "Иванов" not in text
    assert text.count("Экскаватор JCB 3CX") == len(parts)
    # График из 12 платежей: 11 равных и последний с остатком — в каждой части
    # («2 083 337» дважды: остаток после 11-го платежа и сам 12-й платёж).
    assert text.count("2 083 333") == 11 * len(parts)
    assert text.count("2 083 337") == 2 * len(parts)
    assert ("РАСПИСКА" in text) == ("ru" in parts)
    assert ("ТИЛХАТ" in text) == ("uz" in parts)
    # Подписанта кредитора нет ни в одной части: расписку пишет должник.
    for gone in ("в лице", "действующего на основании", "Должность", "М.П.", "номидан", "Лавозими", "М.Ў."):
        assert gone not in text, gone
    if "ru" in parts:
        assert "(далее — «Кредитор») следующий товар: Экскаватор JCB 3CX" in text
        assert "Кредитор: FARID IMPEKS LLC" in text
        assert "общей стоимостью 25 000 000 (двадцать пять миллионов) сум" in text
    else:
        assert "Должник" not in text and "долга" not in text
    if "uz" in parts:
        assert "умумий қиймати 25 000 000 (йигирма беш миллион) сўм".translate(_UZ_DETACHED) in text
        assert "Тошкент ш." in text
    else:
        assert "арздор" not in text and "сўм" not in text
    # Пустой страницы на месте отрезанного разрыва нет: на каждой — текст.
    # (Число страниц зависит от длины графика: блоки подписей держатся
    # вместе и при 12 платежах уходят на третью страницу — это не пустота.)
    assert all(page.strip() for page in pages), f"{doc_type}: пустая страница"
    first = " ".join(pages[0].split())
    assert first.startswith("РАСПИСКА" if "ru" in parts else "ТИЛХАТ")


@pytest.mark.skipif(not HAS_SOFFICE, reason="нет LibreOffice")
def test_render_reports_missing_template():
    import asyncio

    with pytest.raises(ld.DocumentError, match="Шаблон не найден"):
        asyncio.run(ld.render_pdf("raspiska_ru", _ctx(), "/tmp", template_override="/nope.docx"))


# ─── Имя файла и ошибки LibreOffice (без бинаря: мок границы subprocess) ──────


class _FakeSoffice:
    """Подменяет запуск LibreOffice: пишет PDF в --outdir (или не пишет) и
    отдаёт заданный stderr. Всё остальное — шаблон, docxtpl, копирование в
    каталог документов — настоящее."""

    def __init__(self, *, make_pdf=True, stderr=b""):
        self.make_pdf = make_pdf
        self.stderr = stderr

    async def __call__(self, *args, **kwargs):
        from pathlib import Path

        outdir = Path(args[list(args).index("--outdir") + 1])
        fake = self

        class _Proc:
            returncode = 0

            async def communicate(self):
                if fake.make_pdf:
                    (outdir / "document.pdf").write_bytes(b"%PDF-1.4 " + os.urandom(8))
                return b"", fake.stderr

        return _Proc()


def test_second_document_same_day_does_not_overwrite_first(tmp_path, monkeypatch):
    """Имя было {тип}_{ФИО}_{дата}: вторая расписка тому же должнику за день
    затирала первую, и запись первой отдавала чужой PDF."""
    import asyncio

    monkeypatch.setattr(ld.asyncio, "create_subprocess_exec", _FakeSoffice())
    ctx = _ctx()
    first = asyncio.run(ld.render_pdf("raspiska_ru", ctx, tmp_path))
    first_bytes = first.read_bytes()
    second = asyncio.run(ld.render_pdf("raspiska_ru", ctx, tmp_path))

    assert first != second
    assert first.is_file() and second.is_file()
    assert first.read_bytes() == first_bytes, "первый документ перезаписан"
    assert first.name.startswith("raspiska_ru_Иванов_Иван_Иванович_")
    assert len(list(tmp_path.glob("*.pdf"))) == 2


def test_reserve_pdf_path_is_unique_within_one_second(tmp_path, monkeypatch):
    """Два документа в одну секунду (двойное нажатие) — разные файлы."""
    from datetime import datetime as real_dt

    class _Frozen(real_dt):
        @classmethod
        def now(cls, tz=None):
            return real_dt(2026, 9, 15, 10, 0, 0)

    monkeypatch.setattr(ld, "datetime", _Frozen)
    paths = [ld._reserve_pdf_path(tmp_path, "tilxat_uz", "X") for _ in range(3)]
    assert len({p.name for p in paths}) == 3
    assert paths[0].name == "tilxat_uz_X_2026-09-15_100000.pdf"
    assert paths[1].name == "tilxat_uz_X_2026-09-15_100000_2.pdf"


def test_reserve_pdf_path_long_cyrillic_name_fits_255_bytes(tmp_path):
    """Длинное ФИО кириллицей (2 байта на букву) давало ENAMETOOLONG вместо
    документа. Имя режется по байтам, символ пополам не рвётся."""
    long_name = ld._safe_name("Ёлкина-Абдурахмановна " * 30)
    paths = [ld._reserve_pdf_path(tmp_path, "raspiska_ru_uz", long_name) for _ in range(2)]
    for p in paths:
        assert p.is_file()
        assert len(p.name.encode("utf-8")) <= 255
        p.name.encode("utf-8").decode("utf-8")  # целые символы
    assert paths[0].name.startswith("raspiska_ru_uz_Ёлкина")
    assert paths[0] != paths[1]


def test_truncate_utf8_does_not_split_a_character():
    assert ld._truncate_utf8("ЖЖЖ", 5) == "ЖЖ"
    assert ld._truncate_utf8("abc", 10) == "abc"


def test_libreoffice_stderr_goes_to_log_not_to_user(tmp_path, monkeypatch, caplog):
    """stderr LibreOffice уезжал в текст ошибки формы — пути /tmp и английская
    диагностика. Теперь он в логе, человеку — короткая фраза."""
    import asyncio
    import logging

    raw = b"Error: source file could not be loaded /tmp/tmpabc123/document.docx"
    monkeypatch.setattr(
        ld.asyncio, "create_subprocess_exec", _FakeSoffice(make_pdf=False, stderr=raw)
    )
    with caplog.at_level(logging.ERROR, logger=ld.logger.name), pytest.raises(
        ld.DocumentError
    ) as exc:
        asyncio.run(ld.render_pdf("raspiska_ru", _ctx(), tmp_path))

    message = str(exc.value)
    assert "source file" not in message and "/tmp" not in message
    assert "Не удалось сформировать PDF" in message
    assert "source file could not be loaded" in caplog.text
    assert list(tmp_path.glob("*.pdf")) == []
